"""Genera documentos finales con evidencia del cierre geoespacial validado.

Este archivo no transforma datos ni fuentes. Lee los productos derivados para
incluir sus cifras en documentos editables de entrega.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "final"
FACTS = ROOT / "powerbi" / "geospatial" / "hechos_siniestros.csv"
MAP_CORRIDOR = ROOT / "reports" / "mapa_corredor_siniestros_carrera45_50m.png"
DISTANCE_PLOT = ROOT / "reports" / "distribucion_distancias_siniestros_carrera45.png"

NAVY = "17324D"
SLATE = "536779"
LIGHT = "EAF2F4"
GRID = "D9E2E8"


def cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_border(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "000000", size: float = 9.0) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_border(cell)


def setup(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string("202B36")
    normal.paragraph_format.space_after = Pt(5)

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style.paragraph_format.space_after = Pt(7)

    for name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(3)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("MoviManrique45  |  Comuna 3 Manrique  |  Octubre de 2026")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(SLATE)


def title(document: Document, text: str, subtitle: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.add_run(text)
    subtitle_paragraph = document.add_paragraph()
    subtitle_paragraph.paragraph_format.space_after = Pt(14)
    run = subtitle_paragraph.add_run(subtitle)
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(SLATE)


def paragraph(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.line_spacing = 1.06
    p.add_run(text)


def bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        p.add_run(item)


def table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    tbl = document.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.autofit = False
    for idx, header in enumerate(headers):
        cell = tbl.rows[0].cells[idx]
        cell_shading(cell, NAVY)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8.8)
        if widths:
            cell.width = Cm(widths[idx])
    for row_index, values in enumerate(rows):
        cells = tbl.add_row().cells
        for idx, value in enumerate(values):
            if row_index % 2 == 1:
                cell_shading(cells[idx], LIGHT)
            set_cell_text(cells[idx], value, size=8.6)
            if widths:
                cells[idx].width = Cm(widths[idx])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def metrics() -> dict[str, int]:
    facts = pd.read_csv(FACTS)
    categories = facts["categoria_analisis"].value_counts().to_dict()
    return {
        "total": len(facts),
        "principal": int(categories["principal_50m"]),
        "borde": int(categories["sensibilidad_borde"]),
        "revision": int(categories["revision_territorial"]),
        "resto": int(categories["resto_manrique"]),
    }


def add_image(document: Document, image: Path, width_cm: float, caption: str) -> None:
    if not image.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Cm(width_cm))
    caption_p = document.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(4)
    r = caption_p.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.3)
    r.font.color.rgb = RGBColor.from_string(SLATE)


def build_technical(data: dict[str, int]) -> Path:
    doc = Document()
    setup(doc)
    title(doc, "Documento técnico MoviManrique45", "Seguridad vial y accesibilidad en la Carrera 45 de la Comuna 3 Manrique")
    paragraph(doc, "MoviManrique45 organiza evidencia histórica de siniestros para apoyar la lectura territorial de la Carrera 45. El resultado principal es un corredor analítico estricto de 50 metros, limitado por el polígono oficial de Manrique y separado de los casos de borde o revisión.")
    doc.add_heading("Resultados principales", level=1)
    table(doc, ["Indicador", "Resultado"], [
        ["Siniestros nominales de Comuna 3", f"{data['total']:,}".replace(",", ".")],
        ["Corredor territorial estricto de 50 m", f"{data['principal']:,}".replace(",", ".")],
        ["Sensibilidad de borde", str(data["borde"])],
        ["Revisión territorial", str(data["revision"])],
        ["Resto de Manrique", f"{data['resto']:,}".replace(",", ".")],
        ["Eje vial", "26 segmentos y 1.698,67 m"],
        ["Concentraciones exploratorias", "439 celdas y 12 zonas"],
    ], [10.5, 6.0])
    paragraph(doc, "El indicador de 604 no equivale a los 644 registros obtenidos por proximidad geométrica simple de hasta 50 metros. Los 40 casos adicionales cercanos al borde territorial se conservan como sensibilidad y no amplían el indicador oficial.")

    doc.add_page_break()
    doc.add_heading("Problema territorial y objetivo", level=1)
    paragraph(doc, "La seguridad vial y la accesibilidad peatonal requieren evidencia espacial que distinga la cercanía a un eje vial de la pertenencia territorial. El proyecto se concentra exclusivamente en la Comuna 3 Manrique y no incorpora segmentos externos por su atributo nominal cuando su geometría queda fuera del polígono oficial.")
    doc.add_heading("Fuentes y metodología", level=1)
    bullets(doc, [
        "Siniestros viales históricos de Medellín, filtrados reproduciblemente para Comuna 3.",
        "Malla vial y límite oficial de Manrique provenientes de la Alcaldía de Medellín.",
        "Eje de Carrera 45 recortado contra el límite territorial, con 26 segmentos válidos.",
        "Buffer analítico de 50 m recortado territorialmente. La distancia representa proximidad, no causalidad.",
        "Categorización exclusiva: corredor principal, sensibilidad de borde, revisión territorial y resto de Manrique.",
        "Cuadrícula de 100 m para concentraciones históricas exploratorias. No estima riesgo ni significancia estadística.",
    ])
    doc.add_heading("Arquitectura tecnológica", level=1)
    paragraph(doc, "Python, pandas, GeoPandas y Shapely procesan datos locales. GeoPackage y GeoJSON preservan capas geográficas. Un dashboard web estático consume una vista derivada validada y QGIS conserva el proyecto de inspección geográfica. Power BI cuenta con recursos de implementación preparados, pero no con un archivo PBIX o PBIP validado.")

    doc.add_page_break()
    doc.add_heading("Resultados geoespaciales", level=1)
    add_image(doc, MAP_CORRIDOR, 13.4, "Corredor analítico de 50 m, eje definitivo y siniestros seleccionados. Fuente: resultados geoespaciales del proyecto.")
    paragraph(doc, "La delimitación territorial evita interpretar como parte del corredor los registros próximos al eje pero ubicados fuera del polígono oficial de Manrique. Se preservan seis fragmentos menores de un metro para mantener trazabilidad del recorte geométrico.")

    doc.add_page_break()
    doc.add_heading("Dashboard y análisis exploratorio", level=1)
    add_image(doc, DISTANCE_PLOT, 13.4, "Distribución de distancias al eje de Carrera 45. Fuente: análisis de distancias del proyecto.")
    paragraph(doc, "El dashboard local presenta indicadores, evolución anual, distribución por gravedad y tipo, filtros por año, barrio y clasificación, además del geovisor. La aplicación bloquea la exportación si cambian los 8.332 eventos o los 604 IDs principales.")
    doc.add_heading("Alcance analítico e IA", level=1)
    paragraph(doc, "El proyecto no implementa un modelo de aprendizaje automático ni una predicción de siniestros. La grilla de concentraciones constituye análisis geoespacial descriptivo. Faltan variables de exposición, aforos, velocidades observadas, infraestructura detallada y validación prospectiva para justificar un modelo predictivo responsable.")

    doc.add_page_break()
    doc.add_heading("Limitaciones y recomendaciones", level=1)
    bullets(doc, [
        "Los registros históricos no sustituyen inspección de campo, aforos ni auditoría de accesibilidad.",
        "La ubicación válida no garantiza precisión absoluta del geocodificado.",
        "Las zonas exploratorias no son tasas de riesgo, puntos negros oficiales ni evidencia de causalidad.",
        "Los cuatro registros de revisión territorial requieren contraste con fuentes autorizadas antes de reclasificarlos.",
    ])
    paragraph(doc, "Las propuestas preliminares incluyen revisión de cruces peatonales, visibilidad, continuidad de andenes, señalización, iluminación, gestión de velocidad y accesibilidad universal. La evaluación de reductores, señalización, protección de personas mayores y personas con discapacidad requiere diseño autorizado, participación comunitaria y validación de la autoridad competente.")
    paragraph(doc, "Las recomendaciones no acreditan obra ejecutada, aprobación institucional ni diseño de ingeniería. El proyecto entrega una base reproducible para priorizar futuras revisiones técnicas.")
    path = OUT / "MoviManrique45_Documento_Tecnico.docx"
    doc.save(path)
    return path


def build_executive(data: dict[str, int]) -> Path:
    doc = Document()
    setup(doc)
    title(doc, "Resumen ejecutivo MoviManrique45", "Seguridad vial y accesibilidad de la Carrera 45 en la Comuna 3 Manrique")
    paragraph(doc, "MoviManrique45 reúne datos históricos de siniestros y análisis geoespacial para apoyar la priorización técnica de la movilidad y accesibilidad peatonal en la Carrera 45. El prototipo local ofrece indicadores, filtros y un geovisor con resultados trazables.")
    table(doc, ["Resultado", "Valor"], [
        ["Siniestros analizados en Comuna 3", f"{data['total']:,}".replace(",", ".")],
        ["Casos del corredor territorial estricto", str(data["principal"])],
        ["Casos de sensibilidad de borde", str(data["borde"])],
        ["Casos de revisión territorial", str(data["revision"])],
        ["Zonas prioritarias exploratorias", "12"],
    ], [11.0, 5.5])
    doc.add_heading("Qué significa el corredor", level=1)
    paragraph(doc, "El corredor se construyó con un buffer de 50 metros alrededor del eje definitivo de Carrera 45 y se recortó al límite oficial de Manrique. Por ello, los 604 casos principales son diferentes de los 644 eventos que aparecen al medir solo la proximidad geométrica al eje.")
    doc.add_heading("Uso recomendado", level=1)
    paragraph(doc, "El tablero permite revisar años, barrios y categorías territoriales. Las 12 zonas exploratorias orientan revisiones de cruces, visibilidad, señalización, continuidad de andenes, velocidad y accesibilidad. Las decisiones de inversión requieren visita de campo, diseño de ingeniería, participación comunitaria y validación institucional.")
    doc.add_heading("Límites", level=1)
    paragraph(doc, "Los datos describen siniestros históricos. No calculan tasas de riesgo, no atribuyen causalidad a la vía y no predicen accidentes. El proyecto conserva 40 casos de borde y cuatro registros de revisión fuera del indicador principal para mantener transparencia territorial.")
    path = OUT / "MoviManrique45_Resumen_Ejecutivo.docx"
    doc.save(path)
    return path


def build_budget() -> Path:
    doc = Document()
    setup(doc)
    title(doc, "Presupuesto estimado MoviManrique45", "Proyección de implementación integral durante seis meses")
    paragraph(doc, "Las cifras son estimaciones de planeación para un piloto integral. No constituyen cotizaciones, ofertas comerciales ni presupuesto aprobado. El total no incorpora obras físicas, las cuales requieren estudios, diseños, permisos y contratación.")
    doc.add_heading("Estimación de seis meses", level=1)
    table(doc, ["Rubro", "Supuesto de planeación", "COP"], [
        ["Dirección técnica y gestión", "Dedicación parcial durante seis meses", "54.000.000"],
        ["Ingeniería de datos y geoespacial", "Dos perfiles parciales durante seis meses", "96.000.000"],
        ["Diseño y desarrollo de visualización", "Un perfil parcial durante cinco meses", "35.000.000"],
        ["Validación técnica y QA", "Revisión independiente y pruebas", "22.000.000"],
        ["Trabajo de campo y participación", "Aforos, accesibilidad y talleres, sujeto a permisos", "48.000.000"],
        ["Infraestructura y licencias", "Servicios locales o nube, respaldo y herramientas", "18.000.000"],
        ["Contingencia", "10 % de costos anteriores", "27.300.000"],
        ["Total estimado sin obra física", "", "300.300.000"],
    ], [4.4, 8.0, 3.3])
    doc.add_heading("Intervenciones físicas", level=1)
    paragraph(doc, "Las intervenciones físicas forman un rubro separado y no se cuantifican con los datos actuales. La definición de cruces, señalización, accesibilidad, calmado de tráfico u otras obras requiere inspección de campo, estudios de tránsito, diseño autorizado y proceso contractual.")
    doc.add_heading("Condiciones de uso", level=1)
    paragraph(doc, "La estimación sirve para planeación inicial. Cualquier decisión presupuestal debe actualizar salarios, licencias, precios de infraestructura, alcance del trabajo de campo y requerimientos de la autoridad competente.")
    path = OUT / "MoviManrique45_Presupuesto_Seis_Meses.docx"
    doc.save(path)
    return path


def build_manual() -> Path:
    doc = Document()
    setup(doc)
    title(doc, "Manual de demostración MoviManrique45", "Ejecución local del dashboard y verificación de resultados")
    doc.add_heading("Requisitos", level=1)
    paragraph(doc, "Use el entorno virtual incluido o cree uno con las dependencias declaradas. La demostración web no necesita conexión a servicios externos ni acceso a PostgreSQL.")
    doc.add_heading("Procedimiento", level=1)
    table(doc, ["Paso", "Acción"], [
        ["1", "Desde la raíz del proyecto ejecute scripts\\iniciar_demo.ps1."],
        ["2", "Abra http://localhost:8501 en un navegador."],
        ["3", "Revise los indicadores 8.332, 604, 40 y 4."],
        ["4", "Aplique filtros por año, barrio y clasificación."],
        ["5", "Explore el geovisor, sus capas, leyenda e información de zonas."],
        ["6", "Ejecute la suite de pruebas indicada por el manual del repositorio antes de una presentación formal."],
    ], [1.4, 14.3])
    doc.add_heading("Criterios de demostración", level=1)
    paragraph(doc, "El indicador principal corresponde a 604 siniestros dentro del corredor territorial estricto de 50 metros. Los 40 casos de sensibilidad de borde y cuatro de revisión territorial aparecen separados. La proximidad no demuestra causalidad y las zonas exploratorias no son tasas de riesgo.")
    doc.add_heading("Power BI", level=1)
    paragraph(doc, "Los recursos de Power Query, DAX y modelo están listos para importación manual. No presente Power BI como terminado hasta crear y validar un PBIX o PBIP en Power BI Desktop.")
    path = OUT / "MoviManrique45_Manual_Demostracion.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    data = metrics()
    paths = [build_technical(data), build_executive(data), build_budget(), build_manual()]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
