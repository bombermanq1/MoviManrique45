"""Crea versiones v2 editables de los entregables, sin sobrescribir v1.

Las cifras se leen de productos validados y de la vista derivada del módulo de
intervenciones. No transforma fuentes RAW ni capas geoespaciales.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ESTILO = os.environ.get("MOVIMANRIQUE_ESTILO", "v2").strip().lower()
ES_V3 = ESTILO == "v3"
OUT = ROOT / "deliverables" / "final" / ("v3_textos_claros" if ES_V3 else "v2")
SUFIJO_ARCHIVO = "v3_Textos_Claros" if ES_V3 else "v2_Intervenciones"
FACTS = ROOT / "powerbi" / "geospatial" / "hechos_siniestros.csv"
INTERVENCIONES = ROOT / "web" / "data" / "intervenciones.json"
MAPA_CORREDOR = ROOT / "reports" / "mapa_corredor_siniestros_carrera45_50m.png"
MAPA_QGIS = ROOT / "reports" / "geospatial" / "qgis_final_preview.png"

NAVY = "0F172A"
BLUE = "2563EB"
TEAL = "14B8A6"
SLATE = "475569"
PALE = "F1F5F9"
BORDER = "CBD5E1"


def texto(version_v2: str, version_v3: str) -> str:
    """Conserva la versión v2 y permite generar una versión clara v3."""
    return version_v3 if ES_V3 else version_v2


def color(run, hex_value: str) -> None:
    run.font.color.rgb = RGBColor.from_string(hex_value)


def set_font(run, size: float, bold: bool = False, hex_value: str = NAVY) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    color(run, hex_value)


def shade(cell, hex_value: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_value)
    properties.append(shading)


def border(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), BORDER)


def page_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.1)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("000000")
    for name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first = footer.add_run("MoviManrique45  |  Comuna 3 - Manrique  |  Página ")
    set_font(first, 8, hex_value=SLATE)
    page_field(footer)


def heading(doc: Document, value: str, level: int = 1) -> None:
    paragraph = doc.add_heading(value, level=level)
    for run in paragraph.runs:
        set_font(run, 15 if level == 1 else 11.5, bold=True, hex_value="000000")


def paragraph(doc: Document, value: str, *, emphasis: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(value)
    set_font(run, 10.1, bold=emphasis)


def bullets(doc: Document, values: list[str]) -> None:
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(value), 9.6)


def table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    for index, value in enumerate(headers):
        cell = tbl.rows[0].cells[index]
        cell.width = Cm(widths_cm[index])
        shade(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(value), 8.7, bold=True, hex_value="FFFFFF")
        border(cell)
    for row_index, values in enumerate(rows):
        cells = tbl.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), 8.5)
            border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def title_block(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    set_font(p.add_run(title), 25, bold=True, hex_value="000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_font(p.add_run(subtitle), 11, hex_value=SLATE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(texto(
        "VERSIÓN 2  |  MÓDULO DE INTERVENCIONES EXPLORATORIAS  |  OCTUBRE DE 2026",
        "VERSIÓN 3  |  INFORMACIÓN PARA LA REVISIÓN DE CAMPO  |  OCTUBRE DE 2026",
    ))
    set_font(run, 8.5, bold=True, hex_value=BLUE)


def add_image(doc: Document, path: Path, width_cm: float, caption: str) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(5)
    run = caption_p.add_run(caption)
    set_font(run, 8, hex_value=SLATE)
    run.italic = True


def facts() -> tuple[dict[str, int], dict]:
    data = pd.read_csv(FACTS)
    categories = data["categoria_analisis"].value_counts().to_dict()
    interventions = json.loads(INTERVENCIONES.read_text(encoding="utf-8"))
    return (
        {
            "total": len(data),
            "principal": int(categories["principal_50m"]),
            "borde": int(categories["sensibilidad_borde"]),
            "revision": int(categories["revision_territorial"]),
            "resto": int(categories["resto_manrique"]),
        },
        interventions,
    )


def fmt(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def build_technical(metrics: dict[str, int], module: dict) -> Path:
    doc = Document()
    setup(doc)
    title_block(doc, "Documento técnico MoviManrique45", texto(
        "Seguridad vial, accesibilidad y propuestas exploratorias en la Carrera 45",
        "Seguridad vial y accesibilidad en la Carrera 45, Comuna 3 - Manrique",
    ))
    paragraph(doc, texto(
        "MoviManrique45 organiza evidencia histórica de siniestros para apoyar una lectura territorial de la Carrera 45 en la Comuna 3 - Manrique. La versión 2 incorpora fichas interactivas de propuestas, conservando el cierre geoespacial y sus límites de interpretación.",
        "MoviManrique45 reúne los registros históricos de siniestros y la información territorial para entender mejor lo que pasa alrededor de la Carrera 45, en la Comuna 3 - Manrique. Esta versión incorpora fichas para revisar zonas identificadas sin perder la trazabilidad del análisis.",
    ))
    heading(doc, texto("Resultados del cierre", "¿Qué encontramos?"))
    table(doc, ["Indicador", "Resultado"], [
        ["Registros nominales de Comuna 3", fmt(metrics["total"])],
        [texto("Corredor territorial estricto de 50 m", "Dentro del corredor (50 m)"), str(metrics["principal"])],
        [texto("Sensibilidad de borde / revisión", "Cerca del límite / por revisar"), f"{metrics['borde']} / {metrics['revision']}"],
        [texto("Resto de Manrique", "Otros registros de Manrique"), fmt(metrics["resto"])],
        [texto("Eje y productos exploratorios", "Eje y zonas identificadas"), "26 segmentos, 33 partes, 6 fragmentos; 439 celdas y 12 zonas"],
    ], [9.7, 7.0])
    paragraph(doc, texto(
        "El indicador territorial de 604 eventos se mantiene separado de los 644 registros obtenidos por proximidad geométrica simple. La proximidad espacial no demuestra causalidad.",
        "Los 604 registros dentro del corredor se muestran aparte de los 644 que aparecen por proximidad geométrica simple. Estar cerca del eje no demuestra causalidad.",
    ), emphasis=True)
    doc.add_page_break()
    heading(doc, texto("Datos y metodología", "Cómo se construyó el análisis"))
    bullets(doc, [
        texto("Siniestros viales históricos y capas territoriales oficiales de Medellín, procesados localmente.", "Registros históricos de siniestros y capas territoriales oficiales de Medellín, procesados localmente."),
        texto("Eje definitivo de Carrera 45 recortado al límite oficial de Manrique en EPSG:9377.", "Eje de Carrera 45 recortado al límite oficial de Manrique en EPSG:9377."),
        texto("Corredor analítico de 50 m recortado al límite territorial y categorías exclusivas de cierre.", "Corredor de análisis de 50 m limitado al territorio oficial de Manrique."),
        texto("GeoPackage y GeoJSON como productos geoespaciales; dashboard web local como vista derivada.", "Productos geoespaciales y dashboard local para consultar los resultados."),
        texto("IDs congelados, hashes de insumos y pruebas automatizadas para conservar trazabilidad.", "Identificadores, insumos y pruebas automatizadas para conservar la trazabilidad."),
    ])
    add_image(doc, MAPA_CORREDOR, 15.6, "Corredor analítico de 50 m dentro de Manrique. Fuente: productos geoespaciales validados del proyecto.")
    doc.add_page_break()
    heading(doc, texto("Módulo de intervenciones propuestas", "Zonas identificadas y posibles soluciones"))
    paragraph(doc, texto(
        "El dashboard incorpora 12 fichas de zonas exploratorias reales. Cada ficha muestra registros históricos de la celda, tipos y gravedad disponibles, contexto geométrico del eje, condiciones por verificar y un catálogo de medidas candidatas. Todas las propuestas comienzan como Pendiente de inspección.",
        "El dashboard incorpora 12 fichas de zonas identificadas. Cada una muestra los registros históricos de la celda, los tipos y gravedad disponibles, su relación con el eje, lo que falta por revisar y opciones que podrían evaluarse. Todas empiezan como Pendiente de inspección.",
    ))
    top = sorted(module["zonas"], key=lambda item: item["comparacion_exploratoria"]["orden_por_registros_historicos"])[:5]
    table(doc, ["Orden", "Zona", "Registros", texto("Corredor estricto", "Dentro del corredor"), "Tramos que pasan por la zona"], [[
        str(zona["comparacion_exploratoria"]["orden_por_registros_historicos"]),
        zona["zona_id"],
        str(zona["diagnostico_historico"]["registros_historicos"]),
        str(zona["diagnostico_historico"]["corredor_estricto_50m"]),
        ", ".join(map(str, zona["evidencia_geoespacial"]["segmentos_eje_intersectados_objectid"])) or "Sin intersección",
    ] for zona in top], [1.5, 3.0, 2.2, 2.7, 5.0])
    heading(doc, texto("Medidas candidatas y costo", "Posibles soluciones y qué falta por revisar"), level=2)
    paragraph(doc, texto(
        "El catálogo contempla pasos peatonales, calmado y control de velocidad, señalización, accesibilidad universal, rampas y cruces, protección peatonal y visibilidad. El catálogo orienta inspección y diseño; no asigna obras definitivas. El costo por zona permanece pendiente de estudio técnico y presupuesto porque no existen cantidades de obra ni precios unitarios verificables.",
        "El catálogo reúne opciones sobre pasos peatonales, velocidad, señalización, accesibilidad, rampas, cruces, protección peatonal y visibilidad. Sirve para orientar la inspección y el diseño; no asigna obras definitivas. El costo por zona sigue pendiente porque faltan cantidades de obra y precios unitarios verificables.",
    ))
    doc.add_page_break()
    heading(doc, texto("Geovisor y uso de la evidencia", "Cómo leer el geovisor"))
    add_image(doc, MAPA_QGIS, 15.6, "Vista geoespacial validada en QGIS. El proyecto conserva geometrías y fragmentos para trazabilidad.")
    paragraph(doc, texto(
        "El geovisor permite seleccionar una zona desde el mapa, el listado o la tabla comparativa. La selección resalta su geometría y abre la ficha. La comparación ordena solamente el número histórico de registros por celda, por lo que no representa una tasa de riesgo ni una prioridad institucional.",
        "El geovisor permite seleccionar una zona desde el mapa, el listado o la tabla. Al hacerlo, se resalta su ubicación y se abre su ficha. El orden solo muestra el número histórico de registros por celda; no es una tasa de riesgo ni una prioridad institucional.",
    ))
    doc.add_page_break()
    heading(doc, texto("Límites y recomendaciones de uso", "Qué falta por revisar"))
    bullets(doc, [
        texto("Los registros históricos no incluyen exposición, flujos, velocidades observadas ni inventario completo de infraestructura.", "Los registros históricos no incluyen exposición, flujos, velocidades observadas ni un inventario completo de infraestructura."),
        texto("Las celdas prioritarias no sustituyen una auditoría de seguridad vial, accesibilidad o participación comunitaria.", "Las zonas identificadas no sustituyen una auditoría de seguridad vial, accesibilidad o participación comunitaria."),
        texto("La intersección de una celda con el eje aporta contexto espacial, pero no atribuye un siniestro a la Carrera 45.", "Que una zona toque el eje aporta contexto espacial, pero no atribuye un siniestro a la Carrera 45."),
        texto("La evaluación de cualquier obra requiere inspección, diseño, autorización, presupuesto y validación de la entidad competente.", "Cualquier obra requiere inspección, diseño, autorización, presupuesto y validación de la entidad competente."),
        "No existe ML predictivo ni Power BI nativo validado en el alcance actual.",
    ])
    paragraph(doc, texto(
        "La plataforma sirve para revisar evidencia y estructurar una agenda de campo. Las recomendaciones permanecen como hipótesis técnicas hasta contar con información adicional y validación institucional.",
        "La herramienta sirve para revisar la evidencia y organizar una agenda de campo. Las posibles soluciones requieren información adicional y validación institucional antes de cualquier decisión.",
    ))
    path = OUT / f"MoviManrique45_Documento_Tecnico_{SUFIJO_ARCHIVO}.docx"
    doc.save(path)
    return path


def build_executive(metrics: dict[str, int], module: dict) -> Path:
    doc = Document()
    setup(doc)
    title_block(doc, "Resumen ejecutivo MoviManrique45", texto(
        "Versión 2 con módulo de intervenciones exploratorias",
        "Una herramienta local para revisar la Carrera 45",
    ))
    paragraph(doc, texto(
        "MoviManrique45 integra datos históricos de siniestros, delimitación territorial y visualización geoespacial para apoyar la revisión técnica de la Carrera 45 en la Comuna 3 - Manrique.",
        "MoviManrique45 reúne registros históricos de siniestros, límite territorial y mapas para apoyar la revisión de la Carrera 45 en la Comuna 3 - Manrique.",
    ))
    table(doc, ["Resultado", "Valor"], [
        ["Siniestros analizados", fmt(metrics["total"])],
        [texto("Casos del corredor territorial estricto", "Dentro del corredor (50 m)"), str(metrics["principal"])],
        [texto("Casos de borde y revisión", "Cerca del límite y por revisar"), f"{metrics['borde']} y {metrics['revision']}"],
        [texto("Zonas exploratorias con ficha", "Zonas identificadas con ficha"), str(module["metadatos"]["zonas_exploratorias"])],
        ["Estado inicial de propuestas", module["metadatos"]["estado_inicial"]],
    ], [10.8, 5.9])
    heading(doc, texto("Módulo de intervenciones", "Zonas identificadas"))
    paragraph(doc, texto(
        "El dashboard permite seleccionar una de las 12 zonas desde el mapa, una lista o una tabla comparativa. Cada ficha muestra el conteo histórico de la celda, gravedad y tipos disponibles, contexto geométrico y medidas candidatas para evaluación. El módulo no asigna obras, costos específicos ni prioridades oficiales.",
        "El dashboard permite seleccionar una de las 12 zonas desde el mapa, una lista o una tabla. Cada ficha muestra los registros históricos de la celda, gravedad y tipos disponibles, relación con el eje y posibles soluciones para evaluar. El módulo no asigna obras, costos específicos ni prioridades oficiales.",
    ))
    heading(doc, texto("Uso recomendado", "Para qué sirve"))
    paragraph(doc, texto(
        "La plataforma facilita preparar inspecciones de campo sobre cruces, velocidad, señalización, visibilidad, andenes y accesibilidad. Las decisiones requieren levantamiento técnico, participación, diseño y autorización de la entidad competente.",
        "La herramienta ayuda a preparar inspecciones de campo sobre cruces, velocidad, señalización, visibilidad, andenes y accesibilidad. Las decisiones requieren levantamiento técnico, participación, diseño y autorización de la entidad competente.",
    ))
    heading(doc, texto("Límites", "Lo que los datos no permiten concluir"))
    paragraph(doc, "Los datos no miden exposición ni causalidad. El corredor de 50 m constituye un criterio analítico y las zonas exploratorias no son tasas de riesgo. Power BI nativo y ML predictivo continúan fuera del alcance validado.")
    path = OUT / f"MoviManrique45_Resumen_Ejecutivo_{SUFIJO_ARCHIVO}.docx"
    doc.save(path)
    return path


def build_budget() -> Path:
    doc = Document()
    setup(doc)
    title_block(doc, "Presupuesto estimado MoviManrique45", texto(
        "Versión 2 con módulo de intervenciones exploratorias",
        "Planeación inicial para seis meses",
    ))
    paragraph(doc, texto(
        "Las cifras corresponden a planeación inicial para seis meses. No constituyen cotizaciones, presupuesto aprobado ni costos de obra física. La incorporación del módulo de intervenciones usa datos y visualización ya incluidos en el alcance de desarrollo.",
        "Esta es una planeación inicial de seis meses; no es una cotización, un presupuesto aprobado ni el costo de una obra física. El módulo de zonas usa datos y visualización ya incluidos en el alcance de desarrollo.",
    ))
    heading(doc, "Estimación de seis meses")
    table(doc, ["Rubro", "Supuesto", "COP"], [
        ["Dirección técnica y gestión", "Dedicación parcial durante seis meses", "54.000.000"],
        ["Ingeniería de datos y geoespacial", "Dos perfiles parciales durante seis meses", "96.000.000"],
        ["Diseño y desarrollo de visualización", "Dashboard, geovisor y módulo exploratorio", "35.000.000"],
        ["Validación técnica y QA", "Pruebas y revisión independiente", "22.000.000"],
        ["Trabajo de campo y participación", "Aforos, accesibilidad y talleres, sujeto a permisos", "48.000.000"],
        ["Infraestructura y licencias", "Servicios, respaldo y herramientas", "18.000.000"],
        ["Contingencia", "10 % de costos anteriores", "27.300.000"],
        ["Total estimado sin obra física", "", "300.300.000"],
    ], [4.4, 8.0, 3.3])
    heading(doc, texto("Intervenciones físicas", "Obras y costos por definir"))
    paragraph(doc, texto(
        "Los costos por zona permanecen pendientes de estudio técnico y presupuesto. Pasos, rampas, señalización, control de velocidad, accesibilidad u otras obras necesitan cantidades verificadas, precios unitarios, diseños, permisos y contratación.",
        "Los costos por zona siguen pendientes de estudio técnico y presupuesto. Pasos, rampas, señalización, control de velocidad, accesibilidad u otras obras necesitan cantidades verificadas, precios unitarios, diseños, permisos y contratación.",
    ))
    heading(doc, "Condiciones de uso")
    paragraph(doc, "Actualice salarios, licencias, trabajo de campo e infraestructura antes de tomar una decisión presupuestal. Separe las cotizaciones documentadas de las estimaciones orientativas en cualquier versión posterior.")
    path = OUT / f"MoviManrique45_Presupuesto_Seis_Meses_{SUFIJO_ARCHIVO}.docx"
    doc.save(path)
    return path


def build_manual() -> Path:
    doc = Document()
    setup(doc)
    title_block(doc, "Manual de demostración MoviManrique45", texto(
        "Versión 2 del dashboard y módulo de intervenciones",
        "Guía para recorrer el dashboard local",
    ))
    paragraph(doc, texto(
        "La demostración local no requiere PostgreSQL, Internet ni Power BI Desktop. Antes de servir el sitio, el iniciador verifica los insumos del cierre, los 8.332 eventos, las categorías congeladas y exporta las vistas web.",
        "La demostración funciona en el equipo local; no necesita PostgreSQL, internet ni Power BI Desktop. Antes de abrir el sitio, el iniciador verifica los insumos, los 8.332 registros, las categorías validadas y exporta las vistas web.",
    ))
    heading(doc, "Ejecución")
    table(doc, ["Paso", "Acción"], [
        ["1", "Desde la raíz, ejecute .\\scripts\\iniciar_demo.ps1."],
        ["2", "Abra http://127.0.0.1:8501/."],
        ["3", "Revise los indicadores 8.332, 604, 40 y 4."],
        ["4", "Aplique filtros y consulte el geovisor."],
        ["5", texto("Abra Intervenciones propuestas y seleccione una de las 12 zonas.", "Abra Posibles soluciones y seleccione una de las 12 zonas.")],
        ["6", "Confirme el estado Pendiente de inspección, las posibles soluciones y lo que falta por revisar."],
    ], [1.4, 14.3])
    heading(doc, texto("Criterios de demostración", "Qué mostrar en la demostración"))
    paragraph(doc, texto(
        "El módulo muestra evidencia histórica de cada celda, no una obra aprobada. El mapa resalta la zona elegida y la ficha distingue datos observados, condiciones por verificar y propuestas candidatas. El orden comparativo usa solo registros históricos por celda.",
        "El módulo muestra registros históricos de cada zona, no una obra aprobada. El mapa resalta la zona elegida y la ficha separa qué encontramos, qué revisar en campo y las posibles soluciones. El orden usa solo registros históricos por celda.",
    ))
    heading(doc, texto("Validación", "Antes de cerrar la demostración"))
    paragraph(doc, "Ejecute .\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v. No presente Power BI como terminado hasta contar con un PBIX o PBIP validado. No presente las zonas como tasas de riesgo, predicciones o causalidad.")
    path = OUT / f"MoviManrique45_Manual_Demostracion_{SUFIJO_ARCHIVO}.docx"
    doc.save(path)
    return path


def main() -> None:
    expected = [
        OUT / f"MoviManrique45_Documento_Tecnico_{SUFIJO_ARCHIVO}.docx",
        OUT / f"MoviManrique45_Resumen_Ejecutivo_{SUFIJO_ARCHIVO}.docx",
        OUT / f"MoviManrique45_Presupuesto_Seis_Meses_{SUFIJO_ARCHIVO}.docx",
        OUT / f"MoviManrique45_Manual_Demostracion_{SUFIJO_ARCHIVO}.docx",
    ]
    existentes = [ruta for ruta in expected if ruta.exists()]
    if existentes:
        raise FileExistsError(f"No se sobrescriben entregables existentes: {existentes}")
    OUT.mkdir(parents=True, exist_ok=True)
    metrics, module = facts()
    paths = [
        build_technical(metrics, module),
        build_executive(metrics, module),
        build_budget(),
        build_manual(),
    ]
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
